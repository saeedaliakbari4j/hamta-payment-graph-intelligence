import os
import sys
import docx
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding='utf-8')

doc_path = r"D:\project\hamta\FA_From_Transactional_Data_to_Organizational_Intelligence.docx"
doc = docx.Document(doc_path)

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

for i, p in enumerate(doc.paragraphs[:30]):
    runs = p.runs
    sz_list = []
    szCs_list = []
    font_list = []
    for r in runs[:3]:
        rPr = r._r.get_or_add_rPr()
        sz = rPr.find(qn('w:sz'))
        szCs = rPr.find(qn('w:szCs'))
        rFonts = rPr.find(qn('w:rFonts'))
        sz_val = sz.get(qn('w:val')) if sz is not None else '-'
        szCs_val = szCs.get(qn('w:val')) if szCs is not None else '-'
        font_cs = rFonts.get(qn('w:cs')) if rFonts is not None else '-'
        sz_list.append(sz_val)
        szCs_list.append(szCs_val)
        font_list.append(font_cs)
    
    clean_text = p.text[:45].strip().replace('\n', ' ')
    print(f"P{i:02d} [{p.style.name:<12}] text='{clean_text}' | sz={sz_list} | szCs={szCs_list} | font={font_list}")

print("\n--- TABLES CHECK ---")
for ti, table in enumerate(doc.tables):
    print(f"Table {ti}: {len(table.rows)} rows, {len(table.columns)} cols")
    c0 = table.rows[0].cells[0].paragraphs[0]
    r = c0.runs[0] if c0.runs else None
    if r is not None:
        rPr = r._r.get_or_add_rPr()
        sz = rPr.find(qn('w:sz'))
        szCs = rPr.find(qn('w:szCs'))
        sz_val = sz.get(qn('w:val')) if sz is not None else '-'
        szCs_val = szCs.get(qn('w:val')) if szCs is not None else '-'
        print(f"  Header cell (0,0): text='{c0.text}' sz={sz_val} szCs={szCs_val}")
