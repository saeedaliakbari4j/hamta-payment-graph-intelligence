import docx
import sys
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document("FA_From_Transactional_Data_to_Organizational_Intelligence.docx")

print(f"Total paragraphs: {len(doc.paragraphs)}")

issues = []

for pi, p in enumerate(doc.paragraphs):
    p_text = p.text
    if not p_text.strip():
        continue
    
    # Check runs in this paragraph
    runs = p.runs
    run_info = []
    for ri, r in enumerate(runs):
        rpr = r._r.get_or_add_rPr()
        bidi = rpr.find(qn('w:bidi'))
        rtl = rpr.find(qn('w:rtl'))
        rf = rpr.find(qn('w:rFonts'))
        ascii_font = rf.get(qn('w:ascii')) if rf is not None else None
        cs_font = rf.get(qn('w:cs')) if rf is not None else None
        is_latin = (ascii_font == 'Times New Roman' and (bidi is not None and bidi.get(qn('w:val')) == '0'))
        run_info.append({
            'text': r.text,
            'is_latin': is_latin,
            'ascii_font': ascii_font,
            'cs_font': cs_font
        })
    
    # Look for broken math fragments like '|', 'O', 'log', 'C(m)', etc.
    for ri, r in enumerate(run_info):
        t = r['text']
        if any(frag in t for frag in ['|C', 'C(m)', 'log K', 'd_avg', 'd_card', 'S_covisit', 'S_category', 'Var(Y)', 'Top-K']):
            if not r['is_latin']:
                issues.append((pi, f"Non-latin run with math text: '{t}' in para: '{p_text[:60]}...'"))
        # Check if Big-O is split across runs
        if t.strip() in ['O', 'O(', 'O(|', 'O(|C']:
            issues.append((pi, f"Split Big-O detected at run {ri}: '{t}' in para: '{p_text[:60]}...'"))

print(f"Total potential issues detected: {len(issues)}")
for pi, msg in issues[:30]:
    print(f"P{pi}: {msg}")
